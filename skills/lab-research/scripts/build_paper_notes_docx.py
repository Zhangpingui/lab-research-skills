#!/usr/bin/env python3
"""Build a readable DOCX from paper-notes.md and its local figure assets."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from urllib.parse import unquote, urlparse

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
IMAGE_RE = re.compile(r"^\s*!\[([^\]]*)\]\((.+)\)\s*$")
BULLET_RE = re.compile(r"^\s*[-+*]\s+(.+)$")
NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+(.+)$")
TABLE_SEPARATOR_CELL_RE = re.compile(r"^:?-{3,}:?$")
INLINE_RE = re.compile(
    r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|\[[^\]]+\]\([^)]+\))"
)
COMPANION_SECTION_HEADINGS = {
    "最小实验计划",
    "本次 skill 全流程测试结果",
    "skill 测试结果",
    "agent 分工",
    "参考文献与检索入口",
    "构建与视觉 qa",
    "运行清单",
    "边界说明",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _set_run_font(run, name: str, size: float, *, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _configure_document(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size in (
        ("Title", 18),
        ("Heading 1", 15),
        ("Heading 2", 13),
        ("Heading 3", 12),
        ("Heading 4", 11.5),
        ("Heading 5", 11),
        ("Heading 6", 11),
    ):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True

    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    for borders in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(borders)


def _plain_inline_text(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    value = value.replace("**", "").replace("`", "")
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", value)
    return value


def _normalized_heading_label(value: str) -> str:
    text = _plain_inline_text(value).strip()
    text = re.sub(r"^\d+(?:\.\d+)*[.)、．:：-]?\s*", "", text)
    return re.sub(r"\s+", " ", text).casefold()


def _find_companion_sections(lines: list[str]) -> list[str]:
    headings: list[str] = []
    for line in lines:
        match = HEADING_RE.match(line)
        if not match:
            continue
        heading = _plain_inline_text(match.group(2)).strip()
        if _normalized_heading_label(heading) in COMPANION_SECTION_HEADINGS:
            headings.append(heading)
    return headings


def _add_inline_runs(paragraph, text: str, *, size: float = 11) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            _set_run_font(run, "SimSun", size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, "SimSun", size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, "Consolas", max(9, size - 1))
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, "SimSun", size, italic=True)
        else:
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            assert link is not None
            run = paragraph.add_run(f"{link.group(1)} ({link.group(2)})")
            _set_run_font(run, "SimSun", size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_run_font(run, "SimSun", size)


def _split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    placeholder = "\x00PIPE\x00"
    stripped = stripped.replace(r"\|", placeholder)
    return [cell.strip().replace(placeholder, "|") for cell in stripped.split("|")]


def _is_table_separator(line: str, expected_cells: int) -> bool:
    cells = _split_table_row(line)
    return len(cells) == expected_cells and all(
        TABLE_SEPARATOR_CELL_RE.fullmatch(cell.replace(" ", "")) for cell in cells
    )


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _shade_cell(cell, "D9EAF7")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_inline_runs(paragraph, header, size=9.5)
        for run in paragraph.runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            _add_inline_runs(paragraph, value, size=9.5)
    document.add_paragraph()


def _parse_image_target(raw_target: str) -> str:
    target = raw_target.strip()
    if target.startswith("<"):
        closing = target.find(">")
        if closing < 0:
            raise ValueError(f"invalid Markdown image target: {raw_target!r}")
        return target[1:closing]
    match = re.match(r"^(\S+?)(?:\s+[\"'].*[\"'])?$", target)
    if not match:
        raise ValueError(
            "image paths containing spaces must use angle brackets or percent encoding"
        )
    return match.group(1)


def _resolve_local_image(markdown_path: Path, raw_target: str) -> Path:
    target = unquote(_parse_image_target(raw_target))
    parsed = urlparse(target)
    if parsed.scheme or parsed.netloc:
        raise ValueError(f"remote images are not allowed in paper notes: {target}")
    candidate = Path(target)
    if candidate.is_absolute():
        raise ValueError(f"paper-notes images must use relative paths: {target}")
    notes_root = markdown_path.parent.resolve()
    resolved = (notes_root / candidate).resolve()
    try:
        resolved.relative_to(notes_root)
    except ValueError as exc:
        raise ValueError(f"image escapes the paper-notes directory: {target}") from exc
    if not resolved.is_file():
        raise FileNotFoundError(f"referenced image not found: {resolved}")
    return resolved


def _image_width_inches(path: Path, max_width_inches: float) -> float:
    with Image.open(path) as image:
        width_px, height_px = image.size
        dpi = image.info.get("dpi", (150, 150))
        dpi_x = float(dpi[0]) if isinstance(dpi, tuple) and dpi else 150.0
        if dpi_x < 30 or dpi_x > 1200:
            dpi_x = 150.0
    natural_width = max(0.5, width_px / dpi_x)
    width = min(max_width_inches, natural_width)
    if width_px and height_px:
        rendered_height = height_px / width_px * width
        if rendered_height > 8.0:
            width *= 8.0 / rendered_height
    return max(0.5, width)


def _add_image(document: Document, path: Path, alt: str, max_width_inches: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Inches(_image_width_inches(path, max_width_inches)))
    description = alt.strip() or path.name
    inline._inline.docPr.set("descr", description)
    inline._inline.docPr.set("title", description)


def _starts_block(lines: list[str], index: int) -> bool:
    line = lines[index]
    stripped = line.strip()
    if not stripped:
        return True
    if stripped.startswith("```") or HEADING_RE.match(line) or IMAGE_RE.match(line):
        return True
    if BULLET_RE.match(line) or NUMBER_RE.match(line) or stripped.startswith(">"):
        return True
    if re.fullmatch(r"[-*_]{3,}", stripped):
        return True
    if index + 1 < len(lines) and "|" in line:
        headers = _split_table_row(line)
        if len(headers) >= 2 and _is_table_separator(lines[index + 1], len(headers)):
            return True
    return False


def build_document(
    markdown_path: Path,
    output_path: Path,
    *,
    max_image_width_cm: float = 15.5,
    overwrite: bool = False,
    allow_companion_sections: bool = False,
) -> dict[str, object]:
    markdown_path = markdown_path.resolve()
    if not markdown_path.is_file():
        raise FileNotFoundError(markdown_path)
    if markdown_path.name.lower() != "paper-notes.md":
        raise ValueError("input must be named paper-notes.md")
    if output_path.suffix.lower() != ".docx":
        raise ValueError("output must end in .docx")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"output already exists: {output_path}")
    if not 5 <= max_image_width_cm <= 17:
        raise ValueError("max image width must be between 5 and 17 cm")

    lines = markdown_path.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
    companion_sections = _find_companion_sections(lines)
    if companion_sections and not allow_companion_sections:
        joined = ", ".join(companion_sections)
        raise ValueError(
            "reader-facing paper notes contain companion sections that belong in sidecar "
            f"artifacts: {joined}. Use --allow-companion-sections only when the user "
            "explicitly requests them in the Word report."
        )
    document = Document()
    _configure_document(document)
    first_h1 = True
    figure_count = 0
    table_count = 0
    heading_count = 0
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError("unclosed fenced code block")
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(code_lines))
            _set_run_font(run, "Consolas", 9)
            if language:
                paragraph.style = document.styles["Normal"]
            index += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            text = _plain_inline_text(heading.group(2))
            if level == 1 and first_h1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(text)
                _set_run_font(run, "Microsoft YaHei", 18, bold=True)
                first_h1 = False
            else:
                document.add_heading(text, level=min(level, 6))
            heading_count += 1
            index += 1
            continue

        image = IMAGE_RE.match(line)
        if image:
            path = _resolve_local_image(markdown_path, image.group(2))
            _add_image(document, path, image.group(1), max_image_width_cm / 2.54)
            figure_count += 1
            index += 1
            continue

        if index + 1 < len(lines) and "|" in line:
            headers = _split_table_row(line)
            if len(headers) >= 2 and _is_table_separator(lines[index + 1], len(headers)):
                index += 2
                rows: list[list[str]] = []
                while index < len(lines) and "|" in lines[index] and lines[index].strip():
                    row = _split_table_row(lines[index])
                    if len(row) != len(headers):
                        raise ValueError(
                            f"Markdown table row {index + 1} has {len(row)} cells; "
                            f"expected {len(headers)}"
                        )
                    rows.append(row)
                    index += 1
                _add_table(document, headers, rows)
                table_count += 1
                continue

        bullet = BULLET_RE.match(line)
        if bullet:
            while index < len(lines):
                item = BULLET_RE.match(lines[index])
                if not item:
                    break
                paragraph = document.add_paragraph(style="List Bullet")
                _add_inline_runs(paragraph, item.group(1))
                index += 1
            continue

        numbered = NUMBER_RE.match(line)
        if numbered:
            while index < len(lines):
                item = NUMBER_RE.match(lines[index])
                if not item:
                    break
                paragraph = document.add_paragraph(style="List Number")
                _add_inline_runs(paragraph, item.group(1))
                index += 1
            continue

        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(re.sub(r"^\s*>\s?", "", lines[index]))
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.75)
            _add_inline_runs(paragraph, " ".join(quote_lines))
            for run in paragraph.runs:
                run.italic = True
            continue

        if re.fullmatch(r"[-*_]{3,}", stripped):
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and not _starts_block(lines, index):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = document.add_paragraph()
        _add_inline_runs(paragraph, " ".join(paragraph_lines))

    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {
        "schema_version": "paper-notes-docx-v1",
        "source": str(markdown_path),
        "source_sha256": sha256(markdown_path),
        "output": str(output_path),
        "output_sha256": sha256(output_path),
        "headings": heading_count,
        "tables": table_count,
        "figures": figure_count,
        "companion_sections_included": companion_sections,
        "visual_qa": "not_run",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="paper-notes.md")
    parser.add_argument("output", type=Path, help="output DOCX")
    parser.add_argument("--report", type=Path, help="optional JSON build report")
    parser.add_argument("--max-image-width-cm", type=float, default=15.5)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument(
        "--allow-companion-sections",
        action="store_true",
        help="include experiment-plan, search-log, QA, or boundary sections only when explicitly requested",
    )
    args = parser.parse_args()
    try:
        if args.report and args.report.exists() and not args.overwrite:
            raise FileExistsError(f"report already exists: {args.report}")
        report = build_document(
            args.input,
            args.output,
            max_image_width_cm=args.max_image_width_cm,
            overwrite=args.overwrite,
            allow_companion_sections=args.allow_companion_sections,
        )
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(
                json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except (OSError, UnicodeError, ValueError) as error:
        print(json.dumps({"valid": False, "error": str(error)}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
