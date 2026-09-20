"""Build an editable research-material DOCX from a validated JSON package."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re
import sys
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SCHEMA = "lab-research-document-v1"
BLOCK_TYPES = {"paragraph", "bullets", "table", "page_break"}
VERIFICATION_LEVELS = {"metadata_only", "abstract_only", "full_text", "figure_table"}
CITATION_RE = re.compile(r"\[@([A-Za-z][A-Za-z0-9_-]*)\]")
PLACEHOLDER_RE = re.compile(r"\[\[[^\]\r\n]+\]\]")
REFERENCE_ID_RE = re.compile(r"R\d+")


def _nonempty_string(value: Any) -> bool:
    return isinstance(value, str) and bool(value.strip())


def _document_texts(spec: dict[str, Any]) -> Iterable[str]:
    for key in ("title", "subtitle"):
        value = spec.get(key)
        if isinstance(value, str):
            yield value
    for section in spec.get("sections", []):
        if not isinstance(section, dict):
            continue
        heading = section.get("heading")
        if isinstance(heading, str):
            yield heading
        for block in section.get("blocks", []):
            if not isinstance(block, dict):
                continue
            if isinstance(block.get("text"), str):
                yield block["text"]
            for item in block.get("items", []):
                if isinstance(item, str):
                    yield item
            for cell in block.get("headers", []):
                if isinstance(cell, str):
                    yield cell
            for row in block.get("rows", []):
                if isinstance(row, list):
                    for cell in row:
                        if isinstance(cell, str):
                            yield cell


def validate_spec(spec: Any) -> list[str]:
    errors: list[str] = []
    if not isinstance(spec, dict):
        return ["document package must be a JSON object"]
    if spec.get("schema_version") != SCHEMA:
        errors.append(f"schema_version must equal {SCHEMA!r}")
    if not _nonempty_string(spec.get("title")):
        errors.append("title must be a non-empty string")
    if not _nonempty_string(spec.get("document_type")):
        errors.append("document_type must be a non-empty string")
    if not _nonempty_string(spec.get("language")):
        errors.append("language must be a non-empty string")

    sections = spec.get("sections")
    if not isinstance(sections, list) or not sections:
        errors.append("sections must be a non-empty list")
        sections = []
    for section_index, section in enumerate(sections):
        prefix = f"sections[{section_index}]"
        if not isinstance(section, dict):
            errors.append(f"{prefix} must be an object")
            continue
        if not _nonempty_string(section.get("heading")):
            errors.append(f"{prefix}.heading must be a non-empty string")
        if section.get("level") not in {1, 2, 3}:
            errors.append(f"{prefix}.level must be 1, 2, or 3")
        blocks = section.get("blocks")
        if not isinstance(blocks, list) or not blocks:
            errors.append(f"{prefix}.blocks must be a non-empty list")
            continue
        for block_index, block in enumerate(blocks):
            block_path = f"{prefix}.blocks[{block_index}]"
            if not isinstance(block, dict):
                errors.append(f"{block_path} must be an object")
                continue
            block_type = block.get("type")
            if block_type not in BLOCK_TYPES:
                errors.append(f"{block_path}.type must be one of {sorted(BLOCK_TYPES)}")
            elif block_type == "paragraph" and not _nonempty_string(block.get("text")):
                errors.append(f"{block_path}.text must be a non-empty string")
            elif block_type == "bullets":
                items = block.get("items")
                if not isinstance(items, list) or not items or any(
                    not _nonempty_string(item) for item in items
                ):
                    errors.append(f"{block_path}.items must be a non-empty list of strings")
            elif block_type == "table":
                headers = block.get("headers")
                rows = block.get("rows")
                if not isinstance(headers, list) or not headers or any(
                    not _nonempty_string(item) for item in headers
                ):
                    errors.append(f"{block_path}.headers must be a non-empty list of strings")
                    headers = []
                if not isinstance(rows, list) or not rows:
                    errors.append(f"{block_path}.rows must be a non-empty list")
                    rows = []
                for row_index, row in enumerate(rows):
                    if not isinstance(row, list) or len(row) != len(headers):
                        errors.append(
                            f"{block_path}.rows[{row_index}] must contain {len(headers)} cells"
                        )
                    elif any(not isinstance(cell, (str, int, float)) for cell in row):
                        errors.append(f"{block_path}.rows[{row_index}] contains an invalid cell")

    references = spec.get("references", [])
    if not isinstance(references, list):
        errors.append("references must be a list")
        references = []
    reference_ids: list[str] = []
    for index, reference in enumerate(references):
        prefix = f"references[{index}]"
        if not isinstance(reference, dict):
            errors.append(f"{prefix} must be an object")
            continue
        reference_id = reference.get("id")
        if not isinstance(reference_id, str) or not REFERENCE_ID_RE.fullmatch(reference_id):
            errors.append(f"{prefix}.id must match R followed by digits")
        elif reference_id in reference_ids:
            errors.append(f"{prefix}.id duplicates {reference_id!r}")
        else:
            reference_ids.append(reference_id)
        for field in ("authors", "title"):
            if not _nonempty_string(reference.get(field)):
                errors.append(f"{prefix}.{field} must be a non-empty string")
        year = reference.get("year")
        if not isinstance(year, (str, int)) or isinstance(year, bool) or not str(year).strip():
            errors.append(f"{prefix}.year must be a non-empty string or integer")
        if reference.get("verification") not in VERIFICATION_LEVELS:
            errors.append(
                f"{prefix}.verification must be one of {sorted(VERIFICATION_LEVELS)}"
            )

    known_ids = set(reference_ids)
    cited_ids = {match.group(1) for text in _document_texts(spec) for match in CITATION_RE.finditer(text)}
    unknown_ids = sorted(cited_ids - known_ids)
    if unknown_ids:
        errors.append(f"document cites unknown reference IDs: {unknown_ids}")
    return errors


def _replace_citations(text: str, citation_numbers: dict[str, int]) -> str:
    return CITATION_RE.sub(lambda match: f"[{citation_numbers[match.group(1)]}]", text)


def _set_run_font(run, name: str, size: float, *, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    for style_name, size in (("Title", 18), ("Subtitle", 11), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)


def _add_text_paragraph(document: Document, text: str, citation_numbers: dict[str, int], style=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(_replace_citations(text, citation_numbers))
    _set_run_font(run, "SimSun", 10.5)
    return paragraph


def _format_reference(reference: dict[str, Any], number: int) -> str:
    parts = [f"[{number}] {reference['authors']}. {reference['title']}."]
    venue = str(reference.get("venue", "")).strip()
    if venue:
        parts.append(venue.rstrip(".") + ",")
    parts.append(str(reference["year"]).rstrip(".") + ".")
    doi = str(reference.get("doi", "")).strip()
    url = str(reference.get("url", "")).strip()
    if doi:
        parts.append(f"DOI: {doi}.")
    if url:
        parts.append(url)
    return " ".join(parts)


def build_document(
    spec: dict[str, Any],
    output_path: Path,
    *,
    allow_placeholders: bool = False,
    overwrite: bool = False,
) -> dict[str, Any]:
    errors = validate_spec(spec)
    if errors:
        raise ValueError("; ".join(errors))
    placeholders = sorted({match.group(0) for text in _document_texts(spec) for match in PLACEHOLDER_RE.finditer(text)})
    if placeholders and not allow_placeholders:
        raise ValueError(f"unresolved placeholders are not allowed: {placeholders}")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"output already exists: {output_path}")

    references = spec.get("references", [])
    citation_numbers = {reference["id"]: index + 1 for index, reference in enumerate(references)}
    document = Document()
    _configure_styles(document)
    for section in document.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.8)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(_replace_citations(spec["title"], citation_numbers))
    _set_run_font(title_run, "Microsoft YaHei", 18, bold=True)
    subtitle_text = str(spec.get("subtitle", "")).strip()
    if subtitle_text:
        subtitle = document.add_paragraph(style="Subtitle")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.add_run(_replace_citations(subtitle_text, citation_numbers))
        _set_run_font(subtitle_run, "Microsoft YaHei", 11)

    for section in spec["sections"]:
        document.add_heading(
            _replace_citations(section["heading"], citation_numbers), level=section["level"]
        )
        for block in section["blocks"]:
            if block["type"] == "paragraph":
                _add_text_paragraph(document, block["text"], citation_numbers)
            elif block["type"] == "bullets":
                for item in block["items"]:
                    _add_text_paragraph(document, item, citation_numbers, style="List Bullet")
            elif block["type"] == "table":
                table = document.add_table(rows=1, cols=len(block["headers"]))
                table.style = "Table Grid"
                for column_index, header in enumerate(block["headers"]):
                    cell = table.rows[0].cells[column_index]
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    run = cell.paragraphs[0].add_run(_replace_citations(header, citation_numbers))
                    _set_run_font(run, "Microsoft YaHei", 9, bold=True)
                for row in block["rows"]:
                    cells = table.add_row().cells
                    for column_index, value in enumerate(row):
                        cells[column_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        run = cells[column_index].paragraphs[0].add_run(
                            _replace_citations(str(value), citation_numbers)
                        )
                        _set_run_font(run, "SimSun", 9)
            elif block["type"] == "page_break":
                document.add_page_break()

    if references:
        document.add_heading("参考文献", level=1)
        for number, reference in enumerate(references, start=1):
            _add_text_paragraph(document, _format_reference(reference, number), {})

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    digest = hashlib.sha256(output_path.read_bytes()).hexdigest()
    return {
        "schema_version": SCHEMA,
        "output": str(output_path),
        "sha256": digest,
        "sections": len(spec["sections"]),
        "references": len(references),
        "placeholders": placeholders,
        "visual_qa": "not_run",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="UTF-8 JSON document package")
    parser.add_argument("output", type=Path, help="DOCX output path")
    parser.add_argument("--report", type=Path, help="Optional JSON build report")
    parser.add_argument("--allow-placeholders", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    try:
        if args.report and args.report.exists() and not args.overwrite:
            raise FileExistsError(f"report already exists: {args.report}")
        spec = json.loads(args.input.read_text(encoding="utf-8"))
        report = build_document(
            spec,
            args.output,
            allow_placeholders=args.allow_placeholders,
            overwrite=args.overwrite,
        )
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(
                json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
            )
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as error:
        print(json.dumps({"valid": False, "error": str(error)}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
