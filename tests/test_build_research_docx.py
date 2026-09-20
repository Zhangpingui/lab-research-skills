"""Behavior checks for structured research-material Word output."""

import importlib.util
from pathlib import Path
import tempfile
import unittest

from docx import Document


SCRIPT = Path(__file__).resolve().parents[1] / "skills/lab-research/scripts/build_research_docx.py"
SPEC = importlib.util.spec_from_file_location("build_research_docx", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def valid_spec():
    return {
        "schema_version": "lab-research-document-v1",
        "title": "SAR 小目标检测项目材料",
        "document_type": "research_proposal",
        "language": "zh-CN",
        "sections": [
            {
                "heading": "立项依据",
                "level": 1,
                "blocks": [
                    {
                        "type": "paragraph",
                        "text": "该句仅演示结构化引用 [@R01]。",
                    },
                    {"type": "bullets", "items": ["问题", "证据", "边界"]},
                    {
                        "type": "table",
                        "headers": ["主张", "状态"],
                        "rows": [["示例", "已核验"]],
                    },
                ],
            }
        ],
        "references": [
            {
                "id": "R01",
                "authors": "Test Author",
                "title": "Synthetic reference for format testing",
                "venue": "Test Journal",
                "year": "2026",
                "doi": "10.0000/test",
                "verification": "full_text",
            }
        ],
    }


class ResearchDocxTests(unittest.TestCase):
    def test_builds_editable_docx_and_resolves_citation(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "proposal.docx"
            report = MODULE.build_document(valid_spec(), output)
            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("该句仅演示结构化引用 [1]。", text)
            self.assertIn("Synthetic reference for format testing", text)
            self.assertEqual(report["references"], 1)
            self.assertEqual(len(report["sha256"]), 64)

    def test_unknown_citation_is_rejected(self):
        spec = valid_spec()
        spec["sections"][0]["blocks"][0]["text"] = "Unknown [@R99]"
        self.assertTrue(any("unknown reference" in item for item in MODULE.validate_spec(spec)))

    def test_placeholder_requires_explicit_draft_mode(self):
        spec = valid_spec()
        spec["sections"][0]["blocks"][0]["text"] = "[[待补实验条件]]"
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "draft.docx"
            with self.assertRaises(ValueError):
                MODULE.build_document(spec, output)
            report = MODULE.build_document(spec, output, allow_placeholders=True)
            self.assertEqual(report["placeholders"], ["[[待补实验条件]]"])

    def test_existing_output_is_not_overwritten_by_default(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "proposal.docx"
            output.write_bytes(b"existing")
            with self.assertRaises(FileExistsError):
                MODULE.build_document(valid_spec(), output)

    def test_table_width_mismatch_is_rejected(self):
        spec = valid_spec()
        spec["sections"][0]["blocks"][2]["rows"] = [["only one cell"]]
        self.assertTrue(any("must contain 2 cells" in item for item in MODULE.validate_spec(spec)))

    def test_missing_reference_metadata_is_rejected(self):
        spec = valid_spec()
        spec["references"][0]["authors"] = None
        self.assertTrue(any("authors must be" in item for item in MODULE.validate_spec(spec)))


if __name__ == "__main__":
    unittest.main()
