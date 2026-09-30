import importlib.util
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from docx import Document
from docx.oxml.ns import qn


SCRIPT = (
    Path(__file__).resolve().parents[1]
    / "skills"
    / "lab-research"
    / "scripts"
    / "build_paper_notes_docx.py"
)
SPEC = importlib.util.spec_from_file_location("build_paper_notes_docx", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(MODULE)


class PaperNotesDocxTests(unittest.TestCase):
    def _fixture(self, root: Path) -> tuple[Path, Path]:
        assets = root / "paper-notes-assets"
        assets.mkdir()
        image_path = assets / "fig-2-page-4.png"
        Image.new("RGB", (900, 450), "white").save(image_path, dpi=(150, 150))
        notes = root / "paper-notes.md"
        notes.write_text(
            """# DenoDet 论文解读

## 核心意思

作者提出 **多子空间去噪** 方法，并报告了主要实验结果。

- 方法机制
- 证据边界

![原文图 2](paper-notes-assets/fig-2-page-4.png)

来源：原文图 2，PDF 物理页码 4。

| 方法 | 指标 |
|---|---:|
| Baseline | 0.81 |
| DenoDet | 0.84 |
""",
            encoding="utf-8",
        )
        return notes, image_path

    def test_builds_docx_with_local_figure_and_table(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            notes, _ = self._fixture(root)
            output = root / "paper-notes.docx"
            report = MODULE.build_document(notes, output)

            document = Document(output)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("DenoDet 论文解读", text)
            self.assertIn("来源：原文图 2", text)
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(report["figures"], 1)
            self.assertEqual(report["tables"], 1)
            self.assertEqual(report["visual_qa"], "not_run")
            title_properties = document.styles["Title"]._element.get_or_add_pPr()
            self.assertEqual(title_properties.findall(qn("w:pBdr")), [])

    def test_missing_figure_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            notes = root / "paper-notes.md"
            notes.write_text("# Notes\n\n![missing](paper-notes-assets/no.png)\n", encoding="utf-8")
            with self.assertRaises(FileNotFoundError):
                MODULE.build_document(notes, root / "paper-notes.docx")

    def test_remote_figure_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            notes = root / "paper-notes.md"
            notes.write_text("# Notes\n\n![remote](https://example.com/a.png)\n", encoding="utf-8")
            with self.assertRaises(ValueError):
                MODULE.build_document(notes, root / "paper-notes.docx")

    def test_figure_must_stay_inside_notes_directory(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            notes_dir = root / "notes"
            notes_dir.mkdir()
            Image.new("RGB", (20, 20), "white").save(root / "outside.png")
            notes = notes_dir / "paper-notes.md"
            notes.write_text("# Notes\n\n![outside](../outside.png)\n", encoding="utf-8")
            with self.assertRaises(ValueError):
                MODULE.build_document(notes, notes_dir / "paper-notes.docx")

    def test_existing_output_is_not_overwritten_by_default(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            notes, _ = self._fixture(root)
            output = root / "paper-notes.docx"
            output.write_bytes(b"existing")
            with self.assertRaises(FileExistsError):
                MODULE.build_document(notes, output)

    def test_input_must_be_named_paper_notes(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            notes = root / "notes.md"
            notes.write_text("# Notes\n", encoding="utf-8")
            with self.assertRaises(ValueError):
                MODULE.build_document(notes, root / "notes.docx")


if __name__ == "__main__":
    unittest.main()
